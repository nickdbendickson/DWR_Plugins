"""Hazard 2.18.51: the hazard tables against ND Dam Safety Standards REG_05.2026, Section 5.

Builds a small structure inventory (no rasters, no model), writes Habitable_Structures as GPKG /
CSV / XLSX with the plugin's own export, writes the classification summary table, and checks:

  * headers: structure type, depth (ft), velocity (fps), arrival, zone and outcome are reported
    (5.2.1.1, 5.4(c) "flow depth in feet, average flow velocity in feet per second (fps), and
    flood wave arrival time"); roads carry AADT (5.2.1.2), railroads trains a day and passenger
    (5.2.1.3);
  * zone names: Danger_Zone is HIGH DANGER ZONE / JUDGMENT ZONE / LOW DANGER ZONE (Figures 5.1 /
    5.2, pp. 21-22), Danger_Chart names Fig. 5.1 (houses built on foundations), Fig. 5.2
    (mobile homes) or the 1D depth > 2 ft rule (5.2.1.1(d)); JUDGMENT, never JUDGEMENT;
  * colours: the XLSX zone cells are filled #EF9897 / #F9F4A3 / #C4E0AD (Figures 5.1 / 5.2);
  * classes: the ND class is High / Medium / Low (5.2); the ACER TM 11 Table 1 class is
    labelled a reference; the criteria table has the 5.2.1.2(a) private / AADT < 500 row;
  * the Runner / Screening findings table (core/nd_hazard.py) carries danger_zone with the
    standard's names when those plugins are next to this one.

    QT_QPA_PLATFORM=offscreen python3 DamSafety_Hazard/tests/test_hazard_tables_reg05.py
"""
import csv
import importlib.util
import os
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
sys.path.insert(0, ROOT)
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from qgis.core import (QgsApplication, QgsCoordinateReferenceSystem, QgsGeometry,  # noqa: E402
                       QgsPointXY, QgsProject)
_app = QgsApplication.instance() or QgsApplication([], True)
_app.setPrefixPath(os.environ.get("QGIS_PREFIX_PATH", "/home/box/miniforge3/envs/qgis334"), True)
_app.initQgis()

from DamSafety_Hazard import structures as ST, classification as CL, linear as LN  # noqa: E402
from DamSafety_Hazard.helpers import export_layer_to_formats  # noqa: E402
from DamSafety_Hazard.xlsx_zone_fill import zone_fills  # noqa: E402
from DamSafety_Hazard.constants import ND_HIGH, ND_MEDIUM, ND_LOW  # noqa: E402

FILLS = {"HIGH DANGER ZONE": "#EF9897", "JUDGMENT ZONE": "#F9F4A3", "LOW DANGER ZONE": "#C4E0AD",
         "HIGH": "#EF9897", "JUDGMENT": "#F9F4A3", "LOW": "#C4E0AD"}
FAILS = []


def check(what, ok, got=""):
    print("{:4} {}{}".format("ok" if ok else "FAIL", what, "" if ok else "  -> {!r}".format(got)))
    if not ok:
        FAILS.append(what)


def rec(i, zone, acer_type, depth, vel, outcome, one_d=False):
    g = QgsGeometry.fromPolygonXY([[QgsPointXY(i * 100, 0), QgsPointXY(i * 100 + 30, 0),
                                    QgsPointXY(i * 100 + 30, 30), QgsPointXY(i * 100, 30)]])
    return {"id": i, "families": {"NSI"}, "fd_id": str(i), "occupancy": "RES1",
            "fema_category": "Residential", "habitable": "Yes", "acer_type": acer_type,
            "type_basis": "test", "type_conflict": False, "lifeline": False, "persons": 3,
            "persons_basis": "test", "area_ft2": 900.0, "max_depth": depth, "max_vel": vel,
            "arrival": 0.5 if depth else None, "zone_found": zone, "zone_mobile": zone,
            "zone": zone, "outcome": outcome, "one_d": one_d, "category": ST.CAT_DWELLING,
            "sources": [], "footprint": True, "geom": g}


def main():
    crs = QgsCoordinateReferenceSystem("EPSG:26914")
    ctx = QgsProject.instance().transformContext()
    structs = [rec(1, 1, ST.TYPE_HOUSE, 8.0, 6.0, ST.OUT_PROBABLE),
               rec(2, 2, ST.TYPE_MOBILE, 2.2, 1.0, ST.OUT_JUDGMENT),
               rec(3, 3, ST.TYPE_HOUSE, 0.5, 0.5, ST.OUT_ECONOMIC),
               rec(4, None, ST.TYPE_HOUSE, None, None, ST.OUT_DRY),
               rec(5, 3, ST.TYPE_HOUSE, 2.5, 0.2, ST.OUT_PROBABLE, one_d=True)]
    tmp = tempfile.mkdtemp(prefix="hz_tables_")
    try:
        lyr = ST.structures_to_layer(structs, crs, include_dry=True)
        _ok, res, paths = export_layer_to_formats(lyr, tmp, "Habitable_Structures", ctx)
        check("Habitable_Structures written as CSV", bool(paths.get("csv")), res)
        with open(paths["csv"], newline="", encoding="utf-8") as f:
            rows = list(csv.DictReader(f))
        head = list(rows[0].keys())
        for col in ("ACER_Type", "Habitable", "Max_Depth_ft", "Max_Vel_fps", "First_Arrival_hr",
                    "ACER_Zone_Label", "Danger_Zone", "Danger_Chart", "ND_Outcome", "Persons"):
            check("structure table reports {} (5.2.1.1 / 5.4(c))".format(col), col in head, head)
        check("units in the headers are ft and fps (5.4(c))",
              not [h for h in head if h.endswith(("_m", "_ms", "_mps"))], head)
        by = {r["Struct_ID"]: r for r in rows}
        want = {"1": ("HIGH DANGER ZONE", "Fig. 5.1"), "2": ("JUDGMENT ZONE", "Fig. 5.2"),
                "3": ("LOW DANGER ZONE", "Fig. 5.1"), "4": ("", ""),
                "5": ("", "depth > 2 ft")}
        for sid, (zone, chart) in want.items():
            r = by.get(sid, {})
            check("structure {}: Danger_Zone {!r}".format(sid, zone),
                  (r.get("Danger_Zone") or "") == zone, r.get("Danger_Zone"))
            check("structure {}: Danger_Chart names {!r}".format(sid, chart or "nothing"),
                  chart in (r.get("Danger_Chart") or "") and (bool(chart) ==
                                                                bool(r.get("Danger_Chart"))),
                  r.get("Danger_Chart"))
        with open(paths["csv"], encoding="utf-8") as f:
            text = f.read()
        check("no JUDGEMENT spelling in the table", "JUDGEMENT" not in text.upper())
        if paths.get("xlsx"):
            fills = zone_fills(paths["xlsx"])
            for label, hexcol in FILLS.items():
                if label in fills:
                    check("XLSX cell fill of {} is {} (Fig. 5.1 / 5.2)".format(label, hexcol),
                          fills[label] == {hexcol}, fills[label])
            check("XLSX has the three zone names filled",
                  {"HIGH DANGER ZONE", "JUDGMENT ZONE", "LOW DANGER ZONE"} <= set(fills), fills)
            check("XLSX dry cells are not filled", fills.get("DRY", {None}) == {None}, fills)
            try:
                import openpyxl                                    # an independent reader
                ws = openpyxl.load_workbook(paths["xlsx"]).active
                hdr = [c.value for c in ws[1]]
                ci = hdr.index("Danger_Zone") + 1
                got = {ws.cell(r, ci).value: ws.cell(r, ci).fill.fgColor.rgb[-6:].upper()
                       for r in range(2, ws.max_row + 1) if ws.cell(r, ci).value}
                check("openpyxl reads the fills", all(
                    "#" + v == FILLS[k] for k, v in got.items()) and len(got) == 3, got)
            except ImportError:
                print("     (openpyxl not installed: independent XLSX read skipped)")
        else:
            print("NOTE XLSX export unavailable in this QGIS build; fills not checked")

        # roads / rail tables (5.2.1.2 / 5.2.1.3)
        for col in ("AADT", "Max_Depth_ft", "Max_Vel_fps", "First_Arrival_hr", "ND_Class",
                    "ND_Basis"):
            check("road table reports {} (5.2.1.2)".format(col),
                  col in [n for n, _t in LN.ROAD_FIELDS])
        for col in ("Trains_Per_Day", "Passenger", "Max_Depth_ft", "First_Arrival_hr",
                    "ND_Class"):
            check("rail table reports {} (5.2.1.3)".format(col),
                  col in [n for n, _t in LN.RAIL_FIELDS])

        # classification summary (5.2 classes, criteria rows, TM 11 labelled a reference)
        roads = [{"nd_class": ND_LOW, "nd_judgment": False, "name": "a", "length_ft": 100.0,
                  "aadt": 50, "aadt_year": 2024, "bridge": False, "persons": 0.0,
                  "road_class": "Other"}]
        nd = CL.classify_nd(structs[:4], roads, [], roads_evaluated=True, rail_evaluated=True)
        check("ND class is High / Medium / Low (5.2)", nd["cls"] in (ND_HIGH, ND_MEDIUM,
                                                                     ND_LOW), nd["cls"])
        crit = [c["criterion"] for c in nd["criteria"]]
        check("criteria list 5.2.1.2(a) private / AADT < 500 roads",
              any("AADT < 500" in c for c in crit), crit)
        low_row = next(c for c in nd["criteria"] if "AADT < 500" in c["criterion"])
        check("5.2.1.2(a) row result is Low", low_row["result"] == ND_LOW, low_row)
        tm11 = CL.classify_tm11(structs[:4], roads)
        paths = CL.write_summary(tmp, ["x"], nd, tm11, ctx)
        with open(paths["csv"], newline="", encoding="utf-8") as f:
            items = {r["Item"]: r["Value"] for r in csv.DictReader(f)}
        check("summary keeps 'ND classification (screening)' (the Runner reads it)",
              items.get("ND classification (screening)") == nd["cls"], items)
        tm_row = [k for k in items if k.startswith("TM 11 classification")]
        check("summary labels the TM 11 class a reference, not the ND class",
              tm_row and "not the ND class" in tm_row[0], tm_row)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    # Runner / Screening findings table (core/nd_hazard.py)
    for p in ("DamSafety_Runner", "DamSafety_Screening"):
        path = os.path.join(ROOT, p, "core", "nd_hazard.py")
        if not os.path.exists(path):
            continue
        spec = importlib.util.spec_from_file_location("_ndh_tab_" + p, path)
        nh = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = nh
        spec.loader.exec_module(nh)
        check(p + " findings table has danger_zone", "danger_zone" in nh.FINDING_COLUMNS,
              nh.FINDING_COLUMNS)
        check(p + " findings depth / velocity in ft and ft/s",
              "depth_ft" in nh.FINDING_COLUMNS and "vel_fts" in nh.FINDING_COLUMNS)
        f = nh.Finding("High", "5.2.1.1", "structure", "1", "x", zone=nh.ZONE_JUDGMENT)
        check(p + " danger_zone of the judgment zone is 'JUDGMENT ZONE'",
              nh.finding_rows(type("C", (), {"findings": [f]})())[0]["danger_zone"] ==
              "JUDGMENT ZONE")
        check(p + " zone names are the standard's", sorted(nh.ZONE_NAMES.values()) ==
              ["HIGH DANGER ZONE", "JUDGMENT ZONE", "LOW DANGER ZONE"], nh.ZONE_NAMES)
        svg = nh.chart_svg(nh.load_charts().charts["house"], [])
        check(p + " chart legend uses the standard's zone names",
              all(n in svg for n in ("HIGH DANGER ZONE", "JUDGMENT ZONE", "LOW DANGER ZONE")))

    print("\n{} failed".format(len(FAILS)) if FAILS else "\nOK: hazard tables match REG_05.2026")
    return 1 if FAILS else 0


if __name__ == "__main__":
    rc = main()
    sys.stdout.flush()
    os._exit(rc)
